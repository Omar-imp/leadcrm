from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required as dj_login_required
from .models import (Customer, Payment, Category, Product,
                    Order, OrderItem, Shipment, ReturnRequest)
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from functools import wraps
from .models import EcommerceUserPermissions
from .decorators import block_ecommerce_view_only
from django.db.models import Sum, Count, F
from django.utils import timezone
from datetime import timedelta
import csv
import io
import openpyxl
import uuid


LOGIN_URL = 'ecommerce_login'


def login_required(view_func):
    return dj_login_required(view_func, login_url='ecommerce_login')


def ecommerce_login(request):
    if request.user.is_authenticated:
        return redirect('ecommerce_dashboard')
    
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user is not None:
            login(request, user)
            return redirect('ecommerce_dashboard')
        messages.error(request, 'Invalid username or password.')
    return render(request, 'ecommerce/login.html')


def ecommerce_logout(request):
    logout(request)
    return redirect('ecommerce_login')


@login_required
def ecommerce_dashboard(request):
    today = timezone.localdate()
    week_ago = today - timedelta(days=7)
    month_start = today.replace(day=1)
    year_start = today.replace(month=1, day=1)

    orders = Order.objects.exclude(status='cancelled')

    revenue_today = sum(o.total for o in orders.filter(created_at__date=today))
    revenue_week = sum(o.total for o in orders.filter(created_at__date__gte=week_ago))
    revenue_month = sum(o.total for o in orders.filter(created_at__date__gte=month_start))
    revenue_year = sum(o.total for o in orders.filter(created_at__date__gte=year_start))

    total_orders = orders.count()
    aov = (revenue_month / orders.filter(created_at__date__gte=month_start).count()) if orders.filter(created_at__date__gte=month_start).count() else 0

    new_customers = Customer.objects.filter(created_at__date__gte=month_start).count()
    total_customers = Customer.objects.count()

    open_tickets = ReturnRequest.objects.filter(status='pending').count()
    pending_orders = Order.objects.filter(status='pending').count()
    low_stock_products = Product.objects.filter(stock_qty__lte=5).count()

    outstanding_payments = Payment.objects.filter(status__in=['pending', 'partial']).aggregate(total=Sum('amount'))['total'] or 0

    context = {
        'active': 'dashboard',
        'revenue_today': revenue_today,
        'revenue_week': revenue_week,
        'revenue_month': revenue_month,
        'revenue_year': revenue_year,
        'total_orders': total_orders,
        'aov': aov,
        'new_customers': new_customers,
        'total_customers': total_customers,
        'open_tickets': open_tickets,
        'pending_orders': pending_orders,
        'low_stock_products': low_stock_products,
        'outstanding_payments': outstanding_payments,
        'recent_orders': orders.select_related('customer').order_by('-created_at')[:5],
    }
    return render(request, 'ecommerce/dashboard.html', context)


@login_required
def customer_list(request):
    customers = Customer.objects.all().order_by('-created_at')
    context = {
        'customers': customers,
        'active': 'customers',
        'title':'All Customers',
    }
    return render(request, 'ecommerce/customer_list.html', context)


@login_required
@block_ecommerce_view_only
def customer_create(request):
    if request.method == 'POST':
        Customer.objects.create(
            name=request.POST.get('name'),
            email=request.POST.get('email'),
            phone=request.POST.get('phone'),
            dob=request.POST.get('dob') or None,
            gender=request.POST.get('gender', ''),
            billing_address=request.POST.get('billing_address', ''),
            shipping_address=request.POST.get('shipping_address', ''),
            tags=request.POST.get('tags', ''),
        )
        messages.success(request, 'Customer created.')
        return redirect('customer_list')

    context = {
        'active': 'customers',
        'title': 'Add Customer',
    }
    return render(request, 'ecommerce/customer_form.html', context)


@login_required
@block_ecommerce_view_only
def bulk_import_customers(request):
    if request.method == 'POST':
        file = request.FILES.get('import_file')
        if not file:
            messages.error(request, 'Please select a CSV or Excel file.')
            return redirect('customer_create')

        filename = file.name.lower()
        rows = []

        try:
            if filename.endswith('.csv'):
                decoded = file.read().decode('utf-8-sig')
                reader = csv.DictReader(io.StringIO(decoded))
                rows = list(reader)
            elif filename.endswith('.xlsx') or filename.endswith('.xls'):
                wb = openpyxl.load_workbook(file, data_only=True)
                ws = wb.active
                headers = [str(cell.value).strip() if cell.value else '' for cell in ws[1]]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    row_dict = dict(zip(headers, row))
                    rows.append(row_dict)
            else:
                messages.error(request, 'Unsupported file type. Please upload a .csv or .xlsx file.')
                return redirect('customer_create')
        except Exception as e:
            messages.error(request, f'Could not read the file: {e}')
            return redirect('customer_create')

        created_count = 0
        skipped_count = 0
        errors = []

        for i, row in enumerate(rows, start=2):
            normalized = {
                str(k).strip().lower().replace(' ', '_'): v
                for k, v in row.items() if k
            }

            name = str(normalized.get('name', '')).strip()
            email = str(normalized.get('email', '')).strip()
            phone = str(normalized.get('phone') or normalized.get('contact_number') or '').strip()
            gender = str(normalized.get('gender', '')).strip()
            billing_address = str(normalized.get('billing_address') or normalized.get('address') or '').strip()
            shipping_address = str(normalized.get('shipping_address', '')).strip()
            tags = str(normalized.get('tags', '')).strip()

            if not name:
                skipped_count += 1
                errors.append(f'Row {i}: missing name, skipped.')
                continue

            try:
                Customer.objects.create(
                    name=name,
                    email=email or None,
                    phone=phone,
                    gender=gender,
                    billing_address=billing_address,
                    shipping_address=shipping_address,
                    tags=tags,
                )
                created_count += 1
            except Exception as e:
                skipped_count += 1
                errors.append(f'Row {i}: {e}')

        if created_count:
            messages.success(request, f'{created_count} customers imported successfully.')
        if skipped_count:
            messages.warning(request, f'{skipped_count} rows skipped.')
        for err in errors[:10]:
            messages.error(request, err)

        return redirect('customer_list')

    return redirect('customer_create')


@login_required
def product_list(request):
    products = Product.objects.select_related('category').order_by('-created_at')
    context = {
        'products': products,
        'active': 'products',
        'title': 'All Products',
    }
    return render(request, 'ecommerce/product_list.html', context)


@login_required
@block_ecommerce_view_only
def product_create(request):
    if request.method == 'POST':
        category_id = request.POST.get('category')
        category = Category.objects.filter(id=category_id).first() if category_id else None
        Product.objects.create(
            name=request.POST.get('name'),
            sku=request.POST.get('sku'),
            barcode=request.POST.get('barcode', ''),
            category=category,
            description=request.POST.get('description', ''),
            cost_price=request.POST.get('cost_price') or 0,
            sale_price=request.POST.get('sale_price') or 0,
            stock_qty=request.POST.get('stock_qty') or 0,
        )
        messages.success(request, 'Product created.')
        return redirect('product_list')

    context = {
        'active': 'products',
        'title': 'Add Product',
        'categories': Category.objects.all(),
    }
    return render(request, 'ecommerce/product_form.html', context) 


@login_required
def category_list(request):
    categories = Category.objects.select_related('parent').order_by('name')
    context = {
        'categories': categories,
        'active': 'categories',
        'title': 'All Categories',
    }
    return render(request, 'ecommerce/category_list.html', context)


@login_required
@block_ecommerce_view_only
def category_create(request):
    if request.method == 'POST':
        parent_id = request.POST.get('parent')
        parent = Category.objects.filter(id=parent_id).first() if parent_id else None
        Category.objects.create(
            name=request.POST.get('name'),
            parent=parent,
        )
        messages.success(request, 'Category created.')
        return redirect('category_list')

    context = {
        'active': 'categories',
        'title': 'Add Category',
        'categories': Category.objects.all(),
    }
    return render(request, 'ecommerce/category_form.html', context)


@login_required
def order_list(request):
    orders = Order.objects.select_related('customer').order_by('-created_at')
    context = {
        'orders': orders,
        'active': 'orders',
        'title': 'All Orders',
    }
    return render(request, 'ecommerce/order_list.html', context)


@login_required
@block_ecommerce_view_only
def order_create(request):
    if request.method == 'POST':
        customer_id = request.POST.get('customer')
        customer = Customer.objects.filter(id=customer_id).first()
        if not customer:
            messages.error(request, 'Please select a valid customer.')
            return redirect('order_create')

        order = Order.objects.create(
            order_number='ORD-' + uuid.uuid4().hex[:8].upper(),
            customer=customer,
            status=request.POST.get('status', 'pending'),
            discount=request.POST.get('discount') or 0,
            tax=request.POST.get('tax') or 0,
            shipping_cost=request.POST.get('shipping_cost') or 0,
            created_by=request.user,
        )

        product_ids = request.POST.getlist('product')
        quantities = request.POST.getlist('quantity')

        for product_id, qty in zip(product_ids, quantities):
            if not product_id or not qty:
                continue
            product = Product.objects.filter(id=product_id).first()
            if product:
                qty_int = int(qty)
                OrderItem.objects.create(
                    order=order,
                    product=product,
                    quantity=qty_int,
                    unit_price=product.sale_price,
                )
                product.stock_qty = max(0, product.stock_qty - qty_int)
                product.save()

        messages.success(request, f'Order {order.order_number} created.')
        return redirect('order_list')

    context = {
        'active': 'orders',
        'title': 'Create Order',
        'customers': Customer.objects.all(),
        'products': Product.objects.all(),
    }
    return render(request, 'ecommerce/order_form.html', context)


@login_required
def order_detail(request, pk):
    order = get_object_or_404(Order, pk=pk)
    context = {
        'order': order,
        'active': 'orders',
        'title': f'Order {order.order_number}',
    }
    return render(request, 'ecommerce/order_detail.html', context)


@login_required
def low_stock_list(request):
    threshold = 5
    products = Product.objects.filter(stock_qty__lte=threshold).order_by('stock_qty')
    context = {
        'products': products,
        'active': 'inventory',
        'title': 'Low Stock Products',
        'threshold': threshold,
    }
    return render(request, 'ecommerce/low_stock_list.html', context)


@login_required
def payment_list(request):
    payments = Payment.objects.select_related('order', 'order__customer').order_by('-created_at')
    context = {
        'payments': payments,
        'active': 'payments',
        'title': 'All Payments',
    }
    return render(request, 'ecommerce/payment_list.html', context)


@login_required
@block_ecommerce_view_only
def payment_create(request, order_id=None):
    order = get_object_or_404(Order, id=order_id) if order_id else None

    if request.method == 'POST':
        order_pk = request.POST.get('order') or order_id
        order_obj = get_object_or_404(Order, id=order_pk)
        Payment.objects.create(
            order=order_obj,
            method=request.POST.get('method', 'cash'),
            status=request.POST.get('status', 'pending'),
            amount=request.POST.get('amount') or 0,
        )
        messages.success(request, 'Payment recorded.')
        return redirect('order_detail', pk=order_obj.pk)

    context = {
        'active': 'payments',
        'title': 'Record Payment',
        'order': order,
        'orders': Order.objects.all(),
    }
    return render(request, 'ecommerce/payment_form.html', context)


@login_required
def shipment_list(request):
    shipments = Shipment.objects.select_related('order', 'order__customer').order_by('-created_at')
    context = {
        'shipments': shipments,
        'active': 'shipping',
        'title': 'All Shipments',
    }
    return render(request, 'ecommerce/shipment_list.html', context)


@login_required
@block_ecommerce_view_only
def shipment_create(request, order_id=None):
    order = get_object_or_404(Order, id=order_id) if order_id else None

    if request.method == 'POST':
        order_pk = request.POST.get('order') or order_id
        order_obj = get_object_or_404(Order, id=order_pk)
        shipment = Shipment.objects.create(
            order=order_obj,
            tracking_number=request.POST.get('tracking_number', ''),
            carrier=request.POST.get('carrier', ''),
            status=request.POST.get('status', 'pickup'),
        )
        if shipment.status == 'delivered':
            order_obj.status = 'delivered'
        elif shipment.status == 'in_transit':
            order_obj.status = 'shipped'
        order_obj.save()

        messages.success(request, 'Shipment created.')
        return redirect('order_detail', pk=order_obj.pk)

    context = {
        'active': 'shipping',
        'title': 'Create Shipment',
        'order': order,
        'orders': Order.objects.all(),
    }
    return render(request, 'ecommerce/shipment_form.html', context)


@login_required
def return_list(request):
    returns = ReturnRequest.objects.select_related('order', 'order__customer', 'product').order_by('-created_at')
    context = {
        'returns': returns,
        'active': 'support',
        'title': 'Return Requests',
    }
    return render(request, 'ecommerce/return_list.html', context)


@login_required
@block_ecommerce_view_only
def return_create(request, order_id=None):
    order = get_object_or_404(Order, id=order_id) if order_id else None

    if request.method == 'POST':
        order_pk = request.POST.get('order') or order_id
        order_obj = get_object_or_404(Order, id=order_pk)
        product_id = request.POST.get('product')
        product = Product.objects.filter(id=product_id).first() if product_id else None

        ReturnRequest.objects.create(
            order=order_obj,
            product=product,
            reason=request.POST.get('reason', ''),
            status='pending',
        )
        messages.success(request, 'Return request created.')
        return redirect('return_list')

    context = {
        'active': 'support',
        'title': 'Create Return Request',
        'order': order,
        'orders': Order.objects.all(),
        'products': Product.objects.all(),
    }
    return render(request, 'ecommerce/return_form.html', context)


@login_required
@block_ecommerce_view_only
def return_update_status(request, pk):
    ret = get_object_or_404(ReturnRequest, pk=pk)
    if request.method == 'POST':
        new_status = request.POST.get('status')
        if new_status in dict(ReturnRequest.STATUS_CHOICES):
            ret.status = new_status
            ret.save()
            messages.success(request, f'Return status updated to {ret.get_status_display()}.')
    return redirect('return_list')


SECTION_LABELS = {
    'dashboard': 'Dashboard',
    'products': 'Products',
    'categories': 'Categories',
    'customers': 'Customers',
    'orders': 'Orders',
    'inventory': 'Inventory',
    'payments': 'Payments',
    'shipping': 'Shipping',
    'returns': 'Returns',
    'reports': 'Reports',
}


@login_required
def ecommerce_user_list(request):
    users = User.objects.all().order_by('username')
    context = {'users': users, 'active': 'users', 'title': 'Users'}
    return render(request, 'ecommerce/user_list.html', context)


@login_required
def ecommerce_user_create(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username already exists.')
        else:
            user = User.objects.create_user(username=username, email=email, password=password)
            user.first_name = request.POST.get('first_name', '')
            user.last_name = request.POST.get('last_name', '')
            user.save()
            _save_ecommerce_permissions(user, request.POST)
            messages.success(request, f'User {username} created.')
            return redirect('ecommerce_user_list')

    return render(request, 'ecommerce/user_form.html', {
        'active': 'users',
        'title': 'Create User',
        'section_labels': SECTION_LABELS,
    })


@login_required
def ecommerce_user_edit(request, pk):
    target_user = get_object_or_404(User, pk=pk)
    if request.method == 'POST':
        target_user.email = request.POST.get('email')
        target_user.first_name = request.POST.get('first_name')
        target_user.last_name = request.POST.get('last_name')
        target_user.save()
        _save_ecommerce_permissions(target_user, request.POST)
        messages.success(request, 'User updated.')
        return redirect('ecommerce_user_list')

    perms, _ = EcommerceUserPermissions.objects.get_or_create(user=target_user)
    return render(request, 'ecommerce/user_form.html', {
        'active': 'users',
        'title': 'Edit User',
        'edit_user': target_user,
        'section_labels': SECTION_LABELS,
        'current_perms': perms,
    })


def _save_ecommerce_permissions(user, post_data):
    perms, _ = EcommerceUserPermissions.objects.get_or_create(user=user)
    for section in EcommerceUserPermissions.ALL_SECTIONS:
        value = post_data.get(f'perm_{section}', 'none')
        if value in ['full', 'view', 'none']:
            setattr(perms, section, value)
    perms.save()


@login_required
def ecommerce_user_delete(request, pk):
    target_user = get_object_or_404(User, pk=pk)
    if target_user == request.user:
        messages.error(request, "You can't delete your own account.")
    else:
        target_user.delete()
        messages.success(request, 'User deleted.')
    return redirect('ecommerce_user_list')